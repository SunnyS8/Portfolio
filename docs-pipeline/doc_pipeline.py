# -*- coding: utf-8 -*-
"""
doc_pipeline.py — конвейер обработки документов.

Читает папку с документами (PDF, DOCX, XLSX, PNG/JPG-сканы), извлекает текст,
классифицирует тип документа, вытаскивает поля (номер, дата, контрагент, ИНН,
сумма) и выгружает свод в Excel.

Использование:
    python doc_pipeline.py <папка>            — обработать папку
    python doc_pipeline.py <папка> --demo     — сгенерировать тестовые документы и обработать
    python doc_pipeline.py --ocr              — проверить, настроен ли OCR

Зависимости: pdfplumber, python-docx, openpyxl, pandas, Pillow, fpdf2,
            pytesseract + Tesseract OCR (опционально, для сканов).
"""

import re
import sys
import io
import argparse
from pathlib import Path
from datetime import datetime

import pandas as pd
import pdfplumber
import docx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageDraw, ImageFont

# ---------------------------------------------------------------- настройки

TESSERACT_CMD = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
OCR_LANGS = "rus+eng"
SHEET_NAME = "Документы"

# ---------------------------------------------------------------- классификация

# Правила: тип документа -> ключевые фразы (ищутся в тексте без учёта регистра)
DOC_TYPES = {
    "Счёт": ["счёт на оплату", "счёт-фактура", "счёт №", "счёт за", "оплату по счёту"],
    "Договор": ["договор поставки", "договор оказания", "договор №", "стороны заключили договор"],
    "Акт": ["акт оказанных услуг", "акт выполненных работ", "акт №", "подписали настоящий акт"],
    "Накладная": ["товарная накладная", "накладная №", "грузополучатель", "отпустил"],
    "Письмо": ["письмо", "уважаемый", "направляем вам"],
    "Отчёт": ["отчёт", "итоги периода", "показатели"],
}

def classify(text: str) -> str:
    """Определяет тип документа по ключевым фразам. Чем больше совпадений, тем лучше."""
    low = text.lower()
    best, score = "Прочее", 0
    for dtype, keys in DOC_TYPES.items():
        hits = sum(1 for k in keys if k in low)
        if hits > score:
            best, score = dtype, hits
    return best

# ---------------------------------------------------------------- извлечение полей

RE_DATE = re.compile(r"(\d{2}[./-]\d{2}[./-]\d{4})")
RE_INN = re.compile(r"\b(?:ИНН|инн)[:\s]*(\d{10}|\d{12})\b")
RE_NUM = re.compile(r"(?:№|#)\s*([A-Za-zА-Яа-я0-9][\w/-]{0,20})")
RE_SUM = re.compile(r"(?:итого|всего к оплате|к оплате|сумма)\s*[:\s]*([0-9][\d\s]*(?:[.,]\d{1,2})?)", re.I)
RE_COMPANY = re.compile(r"([А-ЯЁ][а-яё\s\-\"«»]{2,60}(?:ООО|АО|ЗАО|ИП)\b|\bООО\s*[\"«][^\"»]{2,60}[\"»])")

def extract_fields(text: str) -> dict:
    """Достаёт из текста стандартные поля документа."""
    fields = {"Номер": "", "Дата": "", "Контрагент": "", "ИНН": "", "Сумма": ""}
    if m := RE_DATE.search(text):
        fields["Дата"] = m.group(1)
    if m := RE_INN.search(text):
        fields["ИНН"] = m.group(1)
    if m := RE_NUM.search(text):
        fields["Номер"] = m.group(1)
    if m := RE_SUM.search(text):
        fields["Сумма"] = m.group(1).replace(" ", "").replace(",", ".")
    if m := RE_COMPANY.search(text):
        fields["Контрагент"] = m.group(1).strip()
    return fields

# ---------------------------------------------------------------- чтение документов

def read_pdf(path: Path) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)

def read_docx(path: Path) -> str:
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)

def read_xlsx(path: Path) -> str:
    """Из таблицы берём текстовые ячейки как «текст» документа."""
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.extend(str(c) for c in row if c is not None)
    return "\n".join(parts)

def read_image(path: Path) -> str:
    """OCR картинки или скана через Tesseract."""
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
    return pytesseract.image_to_string(Image.open(path), lang=OCR_LANGS)

READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".png": read_image,
    ".jpg": read_image,
    ".jpeg": read_image,
    ".tif": read_image,
    ".tiff": read_image,
}

def process_folder(folder: Path) -> pd.DataFrame:
    """Проходит по всем документам папки и собирает таблицу результатов."""
    rows = []
    for path in sorted(folder.iterdir()):
        if path.suffix.lower() not in READERS:
            continue
        if path.name.startswith("свод_"):  # не подхватываем свой же выходной файл
            continue
        try:
            text = READERS[path.suffix.lower()](path)
            fields = extract_fields(text)
            rows.append({
                "Файл": path.name,
                "Тип": classify(text),
                **fields,
                "Знаков текста": len(text),
            })
        except Exception as e:  # noqa: BLE001 — на боевых файлах не роняем конвейер
            rows.append({"Файл": path.name, "Тип": "Ошибка чтения", "Знаков текста": 0,
                         "Номер": "", "Дата": "", "Контрагент": "", "ИНН": "", "Сумма": str(e)[:80]})
    return pd.DataFrame(rows)

# ---------------------------------------------------------------- экспорт в Excel

def export_xlsx(df: pd.DataFrame, out: Path):
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=SHEET_NAME, index=False)
        ws = xw.book[SHEET_NAME]
        head_fill = PatternFill("solid", fgColor="4472C4")
        thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
            cell.alignment = Alignment(vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        widths = {"Файл": 30, "Тип": 14, "Номер": 14, "Дата": 12, "Контрагент": 40, "ИНН": 15, "Сумма": 12}
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(col, 12)
        ws.freeze_panes = "A2"

# ---------------------------------------------------------------- демо-документы

def _pdf(path: Path, lines: list[str]):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("DejaVu", "", r"C:\Windows\Fonts\arial.ttf")
    pdf.set_font("DejaVu", size=12)
    for line in lines:
        pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))

def make_demo(folder: Path):
    """Генерирует тестовый комплект: PDF-счёт, DOCX-акт, XLSX-накладную, PNG-скан письма."""
    folder.mkdir(parents=True, exist_ok=True)

    _pdf(folder / "счёт_1042.pdf", [
        "ООО \"Ромашка\"",
        "Счёт № 1042 от 12.10.2026",
        "На оплату: монтаж оборудования",
        "ИНН 7701234567",
        "Итого к оплате: 145 000,00 руб.",
    ])

    d = docx.Document()
    d.add_paragraph('ООО «Север»')
    d.add_paragraph('Акт оказанных услуг № 88 от 15.10.2026')
    d.add_paragraph('ООО «Ромашка» оказала услуги по настройке серверов.')
    d.add_paragraph('Итого: 62 500 руб. К оплате 62 500,00')
    d.save(folder / "акт_88.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Товарная накладная № 551 от 18.10.2026"])
    ws.append(["Грузополучатель: ИП Иванов И.И."])
    ws.append(["ИНН 771234567890"])
    ws.append(["Наименование", "Кол-во", "Сумма"])
    ws.append(["Кабель ВВГ 3x2,5", 120, "18 400,00"])
    ws.append(["Всего к оплате", "", "18 400,00"])
    wb.save(folder / "накладная_551.xlsx")

    img = Image.new("RGB", (900, 400), "white")
    dr = ImageDraw.Draw(img)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 22)
    for i, line in enumerate([
        "ООО «Свет»",
        "Письмо № 7 от 20.10.2026",
        "Уважаемые коллеги! Направляем вам",
        "отчёт по итогам квартала.",
    ]):
        dr.text((40, 40 + i * 60), line, fill="black", font=font)
    img.save(folder / "скан_письмо_7.png")

    print(f"Демо-документы созданы в: {folder}")

# ---------------------------------------------------------------- точка входа

def main():
    ap = argparse.ArgumentParser(description="Конвейер обработки документов")
    ap.add_argument("folder", nargs="?", help="папка с документами")
    ap.add_argument("--demo", action="store_true", help="сгенерировать демо-документы и обработать")
    ap.add_argument("--ocr", action="store_true", help="проверить доступность OCR")
    args = ap.parse_args()

    if args.ocr:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
        try:
            print("OCR доступен, версия:", pytesseract.get_tesseract_version())
            print("Языки:", ", ".join(pytesseract.get_languages(config="")))
        except Exception as e:  # noqa: BLE001
            print("OCR не настроен:", e)
        return

    if not args.folder:
        ap.print_help()
        return

    folder = Path(args.folder)
    if args.demo:
        make_demo(folder)

    df = process_folder(folder)
    if df.empty:
        print("Документов не найдено. Форматы: pdf, docx, xlsx, png, jpg, tif")
        return

    out = folder / "свод_документов.xlsx"
    export_xlsx(df, out)

    print(f"\nОбработано: {len(df)} док. Результат: {out}\n")
    for _, r in df.iterrows():
        print(f"  {r['Файл']:<26} {r['Тип']:<12} №{r['Номер']:<8} {r['Дата']:<11} {r['Сумма']:<12}")

if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    main()
